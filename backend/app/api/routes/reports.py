import hashlib
import html
import json
from datetime import datetime, timezone
from pathlib import Path
from uuid import UUID, uuid4

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.db import get_db
from app.models import AnalysisResult, AnalysisRun, Finding, Report as ReportModel, ReportArtifact as ArtifactModel, ResearchCase
from app.schemas.reports import (
    CreateReportExportRequest, CreateReportRequest, ReportArtifact, ReportDetail,
    ReportListResponse, ReportSection,
)

router = APIRouter()


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def parse_uuid(value: str) -> str:
    try:
        return str(UUID(value))
    except ValueError as exc:
        raise HTTPException(status_code=404, detail="Report resource not found.") from exc


def get_case(db: Session, case_id: str) -> ResearchCase:
    case = db.get(ResearchCase, parse_uuid(case_id))
    if case is None:
        raise HTTPException(status_code=404, detail="Research case not found.")
    return case


def get_report(db: Session, case_id: str, report_id: str) -> ReportModel:
    case = get_case(db, case_id)
    report = db.get(ReportModel, parse_uuid(report_id))
    if report is None or report.case_id != case.id:
        raise HTTPException(status_code=404, detail="Report not found.")
    return report


def serialize_report(db: Session, item: ReportModel) -> ReportDetail:
    snapshot = json.loads(item.snapshot_json)
    artifacts = list(db.scalars(select(ArtifactModel).where(ArtifactModel.report_id == item.id).order_by(ArtifactModel.created_at)))
    return ReportDetail(
        report_id=item.id, case_id=item.case_id, analysis_run_id=item.analysis_run_id,
        generated_at=datetime.fromisoformat(item.generated_at.replace("Z", "+00:00")),
        template_version=item.template_version, report_snapshot_hash=item.report_snapshot_hash,
        title=item.title, issue_display_name=item.issue_display_name,
        information_cutoff_at=datetime.fromisoformat(item.information_cutoff_at.replace("Z", "+00:00")),
        sections=[ReportSection.model_validate(section) for section in snapshot["sections"]],
        warnings=snapshot["warnings"], export_artifacts=[
            ReportArtifact(
                artifact_id=a.id, report_id=a.report_id, format=a.format, sha256=a.sha256,
                size_bytes=a.size_bytes, created_at=datetime.fromisoformat(a.created_at.replace("Z", "+00:00")),
                download_url=f"/api/v1/cases/{item.case_id}/reports/{item.id}/exports/{a.id}/download",
            ) for a in artifacts
        ],
    )


@router.get("/cases/{case_id}/reports", response_model=ReportListResponse)
def list_reports(case_id: str, db: Session = Depends(get_db)) -> ReportListResponse:
    case = get_case(db, case_id)
    rows = list(db.scalars(select(ReportModel).where(ReportModel.case_id == case.id).order_by(ReportModel.generated_at.desc())))
    return ReportListResponse(items=[serialize_report(db, row) for row in rows], total=len(rows))


@router.post("/cases/{case_id}/reports", response_model=ReportDetail, status_code=201)
def create_report(case_id: str, payload: CreateReportRequest, db: Session = Depends(get_db)) -> ReportDetail:
    case = get_case(db, case_id)
    run = db.get(AnalysisRun, parse_uuid(payload.analysis_run_id))
    if run is None or run.case_id != case.id:
        raise HTTPException(status_code=422, detail="The analysis run must belong to this research case.")
    snapshot_input = json.loads(run.input_snapshot_json)
    issue = snapshot_input["issue"]
    results = list(db.scalars(select(AnalysisResult).where(AnalysisResult.analysis_run_id == run.id)))
    findings = list(db.scalars(select(Finding).where(Finding.analysis_run_id == run.id)))
    sections = [
        {
            "sectionCode": "issue_identity",
            "heading": "Confirmed IPO identity",
            "bodyMarkdown": (
                f"**Issuer:** {issue['issuerLegalName']}\n\n**Issue:** {issue['displayName']}\n\n"
                f"**Market segment:** {issue['marketSegment']}\n\n"
                f"**Open / close / listing:** {issue.get('openDate') or 'Unknown'} / {issue.get('closeDate') or 'Unknown'} / {issue.get('listingDate') or 'Unknown'}\n\n"
                f"**Price band:** {issue.get('priceBandLow') or 'Unknown'} – {issue.get('priceBandHigh') or 'Unknown'}"
            ),
            "relatedEvidenceIds": [], "relatedResultIds": [], "relatedFindingIds": [], "statusNote": "Identity was manually confirmed against reviewed case evidence.",
        },
        {
            "sectionCode": "scenario_results",
            "heading": "Listing-day gross-return scenarios",
            "bodyMarkdown": "\n\n".join(
                f"**{result.indicator_code} — {result.display_value or result.value_decimal or 'Unknown'}**\n\n{result.explanation}\n\n*Limitation:* {result.reason or 'No additional note.'}"
                for result in results
            ) or "No eligible calculations were saved in this run.",
            "relatedEvidenceIds": [], "relatedResultIds": [result.id for result in results],
            "relatedFindingIds": [], "statusNote": "These are user-defined scenarios, not realized market outcomes or forecasts.",
        },
        {
            "sectionCode": "evidence_coverage",
            "heading": "Evidence and input coverage",
            "bodyMarkdown": (
                f"Evidence items in the snapshot: {len(snapshot_input['evidence'])}.\n\n"
                f"Manual facts in the snapshot: {len(snapshot_input['facts'])}.\n\n"
                f"Market observations in the snapshot: {len(snapshot_input['observations'])}.\n\n"
                f"Saved scenarios in the snapshot: {len(snapshot_input['scenarios'])}."
            ),
            "relatedEvidenceIds": [item["evidenceId"] for item in snapshot_input["evidence"]],
            "relatedResultIds": [], "relatedFindingIds": [item.id for item in findings],
            "statusNote": "Inputs are preserved as a point-in-time snapshot; review statuses are not inferred.",
        },
        {
            "sectionCode": "findings_and_limitations",
            "heading": "Findings and limitations",
            "bodyMarkdown": "\n\n".join(f"**{finding.title}**\n\n{finding.explanation}\n\n*Limitation:* {finding.limitation_note or 'None recorded.'}" for finding in findings),
            "relatedEvidenceIds": [], "relatedResultIds": [result.id for result in results],
            "relatedFindingIds": [finding.id for finding in findings],
            "statusNote": "Only the enabled gross-scenario formula ran. Financial quality, valuation, demand, GMP trend, cross-signal checks and net returns were not calculated.",
        },
    ]
    warnings = [
        "This report is based on a partial analysis run.",
        "No external source connector is enabled; source data must be supplied and reviewed manually.",
        "Net return is unknown because brokerage, taxes, slippage and other costs are not modelled.",
        "Scenario prices are assumptions, not forecasts or guaranteed outcomes.",
    ]
    generated = now_iso()
    snapshot = {
        "reportId": str(uuid4()), "caseId": case.id, "analysisRunId": run.id, "generatedAt": generated,
        "templateVersion": "p0-report-v1", "title": f"{issue['displayName']} — IPO Forensic Research Report",
        "issueDisplayName": issue["displayName"], "informationCutoffAt": run.information_cutoff_at,
        "sections": sections, "warnings": warnings,
        "sourceAnalysisRunHash": run.input_snapshot_hash,
    }
    canonical = json.dumps(snapshot, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    report_id = snapshot["reportId"]
    report = ReportModel(
        id=report_id, case_id=case.id, analysis_run_id=run.id, generated_at=generated,
        template_version="p0-report-v1", report_snapshot_hash=hashlib.sha256(canonical.encode("utf-8")).hexdigest(),
        title=snapshot["title"], issue_display_name=issue["displayName"],
        information_cutoff_at=run.information_cutoff_at, snapshot_json=canonical,
    )
    db.add(report)
    case.latest_report_id = report.id
    case.last_activity_at = generated
    case.updated_at = generated
    db.commit()
    db.refresh(report)
    return serialize_report(db, report)


@router.get("/cases/{case_id}/reports/{report_id}", response_model=ReportDetail)
def get_report_detail(case_id: str, report_id: str, db: Session = Depends(get_db)) -> ReportDetail:
    return serialize_report(db, get_report(db, case_id, report_id))


@router.post("/cases/{case_id}/reports/{report_id}/exports", response_model=ReportArtifact, status_code=201)
def export_report(
    case_id: str, report_id: str, payload: CreateReportExportRequest, db: Session = Depends(get_db)
) -> ReportArtifact:
    report = get_report(db, case_id, report_id)
    snapshot = json.loads(report.snapshot_json)
    artifact_id = str(uuid4())
    generated = now_iso()
    base = get_settings().data_dir.resolve()
    relative = Path("reports") / report.case_id / report.id / "exports" / f"{artifact_id}.{payload.format}"
    target = (base / relative).resolve()
    if base not in target.parents:
        raise HTTPException(status_code=400, detail="Unsafe export path.")
    target.parent.mkdir(parents=True, exist_ok=True)
    if payload.format == "html":
        parts = [
            "<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width, initial-scale=1'>",
            f"<title>{html.escape(report.title)}</title><style>body{{font:16px/1.6 system-ui,sans-serif;max-width:900px;margin:40px auto;padding:0 20px;color:#20251e}}h1,h2{{line-height:1.2}}.warning{{padding:12px;background:#fff4e5;margin:8px 0}}pre{{white-space:pre-wrap}}</style></head><body>",
            f"<h1>{html.escape(report.title)}</h1><p>Generated {html.escape(report.generated_at)} · Information cutoff {html.escape(report.information_cutoff_at)}</p>",
        ]
        for warning in snapshot["warnings"]:
            parts.append(f"<div class='warning'>{html.escape(warning)}</div>")
        for section in snapshot["sections"]:
            parts.append(f"<section><h2>{html.escape(section['heading'])}</h2><pre>{html.escape(section['bodyMarkdown'])}</pre></section>")
        parts.append("</body></html>")
        content = "".join(parts).encode("utf-8")
        target.write_bytes(content)
    else:
        from docx import Document
        document = Document()
        document.add_heading(report.title, level=0)
        document.add_paragraph(f"Generated: {report.generated_at}")
        document.add_paragraph(f"Information cutoff: {report.information_cutoff_at}")
        for warning in snapshot["warnings"]:
            document.add_paragraph(warning, style="List Bullet")
        for section in snapshot["sections"]:
            document.add_heading(section["heading"], level=1)
            for block in section["bodyMarkdown"].split("\n\n"):
                document.add_paragraph(block.replace("**", "").replace("*", ""))
            if section.get("statusNote"):
                document.add_paragraph(section["statusNote"], style="Caption")
        document.save(target)
        content = target.read_bytes()
    digest = hashlib.sha256(content).hexdigest()
    artifact = ArtifactModel(
        id=artifact_id, report_id=report.id, format=payload.format, sha256=digest,
        size_bytes=len(content), created_at=generated, local_relative_path=relative.as_posix(),
    )
    db.add(artifact)
    db.commit()
    return ReportArtifact(
        artifact_id=artifact.id, report_id=report.id, format=artifact.format, sha256=artifact.sha256,
        size_bytes=artifact.size_bytes, created_at=datetime.fromisoformat(generated.replace("Z", "+00:00")),
        download_url=f"/api/v1/cases/{report.case_id}/reports/{report.id}/exports/{artifact.id}/download",
    )


@router.get("/cases/{case_id}/reports/{report_id}/exports/{artifact_id}/download")
def download_report_export(case_id: str, report_id: str, artifact_id: str, db: Session = Depends(get_db)) -> FileResponse:
    report = get_report(db, case_id, report_id)
    artifact = db.get(ArtifactModel, parse_uuid(artifact_id))
    if artifact is None or artifact.report_id != report.id:
        raise HTTPException(status_code=404, detail="Report export not found.")
    base = get_settings().data_dir.resolve()
    path = (base / artifact.local_relative_path).resolve()
    if base not in path.parents or not path.is_file():
        raise HTTPException(status_code=404, detail="Report export file is unavailable.")
    media_type = "text/html" if artifact.format == "html" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return FileResponse(path, media_type=media_type, filename=f"{report.id}.{artifact.format}")
